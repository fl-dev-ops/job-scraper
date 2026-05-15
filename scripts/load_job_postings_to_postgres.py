# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "openpyxl>=3.1.5",
#   "psycopg[binary]>=3.2.3",
#   "python-dotenv>=1.0.1",
# ]
# ///

"""Load job posting Excel exports into Postgres.

Usage:
    DATABASE_URL="postgresql://..." uv run scripts/load_job_postings_to_postgres.py --site naukri
    uv run scripts/load_job_postings_to_postgres.py --site linkedin --excel-path data/exports/linkedin_jobs.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import psycopg
from dotenv import load_dotenv
from psycopg import sql

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL_PATH = PROJECT_ROOT / "data" / "exports" / "naukri_jobs.xlsx"
DEFAULT_TABLE_NAME = "job_postings"
VALID_SITES = {"naukri", "linkedin", "indeed"}

EXPECTED_HEADERS = [
    "Job title",
    "Company name",
    "Role type",
    "Location",
    "Role summary",
    "Key responsibilities",
    "Required skills",
    "Focus rounds",
    "Focus round pattern",
    "salary_inr_per_year_min",
    "salary_inr_per_year_max",
    "experience_min_years",
    "experience_max_years",
]

COLUMN_MAP = {
    "Job title": "job_title",
    "Company name": "company_name",
    "Role type": "role_type",
    "Location": "location",
    "Role summary": "role_summary",
    "Key responsibilities": "key_responsibilities",
    "Required skills": "required_skills",
    "Focus rounds": "focus_rounds",
    "Focus round pattern": "focus_round_pattern",
    "salary_inr_per_year_min": "salary_inr_per_year_min",
    "salary_inr_per_year_max": "salary_inr_per_year_max",
    "experience_min_years": "experience_min_years",
    "experience_max_years": "experience_max_years",
}

INTEGER_FIELDS = {
    "salary_inr_per_year_min",
    "salary_inr_per_year_max",
    "experience_min_years",
    "experience_max_years",
}


@dataclass(frozen=True)
class LoadedRow:
    values: dict[str, Any]
    source_row_number: int


def _valid_identifier(name: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name):
        raise ValueError(f"Invalid SQL identifier: {name!r}")
    return name


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _clean_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value)
    text = str(value).strip().replace(",", "")
    return int(float(text)) if text else None


def _row_hash(values: dict[str, Any]) -> str:
    payload = {key: values.get(key) for key in ["site", *COLUMN_MAP.values()]}
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=False, default=str)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def read_workbook(path: Path) -> tuple[str, list[LoadedRow]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Workbook has no rows: {path}")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    if headers != EXPECTED_HEADERS:
        raise ValueError(
            "Unexpected workbook headers.\n"
            f"Expected: {EXPECTED_HEADERS}\n"
            f"Found:    {headers}"
        )

    loaded: list[LoadedRow] = []
    for index, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue

        values: dict[str, Any] = {}
        for header, value in zip(EXPECTED_HEADERS, row, strict=True):
            column = COLUMN_MAP[header]
            values[column] = _clean_int(value) if column in INTEGER_FIELDS else _clean_text(value)

        if not values["job_title"]:
            raise ValueError(f"Missing required title at source row {index}")

        loaded.append(LoadedRow(values=values, source_row_number=index))

    return worksheet.title, loaded


def create_table(conn: psycopg.Connection[Any], table_name: str) -> None:
    table = sql.Identifier(table_name)
    conn.execute(
        sql.SQL(
            """
            CREATE TABLE IF NOT EXISTS {table} (
                id BIGSERIAL PRIMARY KEY,
                site TEXT NOT NULL CHECK (site IN ('naukri', 'linkedin', 'indeed')),
                job_title TEXT NOT NULL,
                company_name TEXT,
                role_type TEXT,
                location TEXT,
                role_summary TEXT,
                key_responsibilities TEXT,
                required_skills TEXT,
                focus_rounds TEXT,
                focus_round_pattern TEXT,
                salary_inr_per_year_min INTEGER CHECK (salary_inr_per_year_min IS NULL OR salary_inr_per_year_min >= 0),
                salary_inr_per_year_max INTEGER CHECK (salary_inr_per_year_max IS NULL OR salary_inr_per_year_max >= 0),
                experience_min_years INTEGER CHECK (experience_min_years IS NULL OR experience_min_years >= 0),
                experience_max_years INTEGER CHECK (experience_max_years IS NULL OR experience_max_years >= 0),
                row_hash TEXT NOT NULL UNIQUE,
                loaded_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
            )
            """
        ).format(table=table)
    )
    for column in ("site", "role_type", "company_name", "location"):
        conn.execute(
            sql.SQL("CREATE INDEX IF NOT EXISTS {index} ON {table} ({column})").format(
                index=sql.Identifier(f"{table_name}_{column}_idx"),
                table=table,
                column=sql.Identifier(column),
            )
        )


def upsert_rows(
    conn: psycopg.Connection[Any],
    table_name: str,
    rows: list[LoadedRow],
    *,
    site: str,
) -> None:
    table = sql.Identifier(table_name)
    insert_columns = [
        "site",
        *COLUMN_MAP.values(),
        "row_hash",
    ]
    update_columns = [column for column in insert_columns if column != "row_hash"]
    query = sql.SQL(
        """
        INSERT INTO {table} ({columns})
        VALUES ({placeholders})
        ON CONFLICT (row_hash) DO UPDATE SET
            {updates},
            updated_at = now()
        """
    ).format(
        table=table,
        columns=sql.SQL(", ").join(sql.Identifier(column) for column in insert_columns),
        placeholders=sql.SQL(", ").join(sql.Placeholder() for _ in insert_columns),
        updates=sql.SQL(", ").join(
            sql.SQL("{column} = EXCLUDED.{column}").format(column=sql.Identifier(column))
            for column in update_columns
        ),
    )

    params = []
    for row in rows:
        values = dict(row.values)
        values["site"] = site
        values["row_hash"] = _row_hash(values)
        params.append(tuple(values[column] for column in insert_columns))

    with conn.cursor() as cursor:
        cursor.executemany(query, params)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--site", required=True, choices=sorted(VALID_SITES))
    parser.add_argument("--excel-path", type=Path, default=DEFAULT_EXCEL_PATH)
    parser.add_argument("--table-name", default=DEFAULT_TABLE_NAME)
    parser.add_argument("--database-url", default=None)
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def main() -> None:
    load_dotenv()
    args = parse_args()
    table_name = _valid_identifier(args.table_name)
    database_url = args.database_url or os.getenv("DATABASE_URL")
    excel_path = args.excel_path.expanduser().resolve()

    source_sheet, rows = read_workbook(excel_path)
    print(f"Read {len(rows)} rows from {excel_path} ({source_sheet!r}).")
    print(f"Target table: {table_name}")
    print(f"Site: {args.site}")

    if args.dry_run:
        return

    if not database_url:
        raise SystemExit("Set DATABASE_URL or pass --database-url.")

    with psycopg.connect(database_url) as conn:
        create_table(conn, table_name)
        upsert_rows(
            conn,
            table_name,
            rows,
            site=args.site,
        )
        conn.commit()

    print(f"Upserted {len(rows)} rows into {table_name}.")


if __name__ == "__main__":
    main()
